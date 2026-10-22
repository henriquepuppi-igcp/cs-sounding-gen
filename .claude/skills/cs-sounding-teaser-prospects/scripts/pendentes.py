"""Lista os prospects da base de CS que ainda nao tem sounding teaser.

Uso:
  python3 pendentes.py <base> <pastas.json> <estado.json|-> [--limite 25] [--hoje AAAA-MM-DD] [--criar-corte 15]

- base: o _base_prospects.csv (export diario), o .xlsx da base (Box Drive) ou o .txt/.json devolvido por mcp__Box__get_file_content.
- pastas.json: lista com os nomes das pastas em Sounding Teasers (ex.: ["5334-Cobuccio", ...]).
- estado.json: o _estado.json da raiz de Sounding Teasers, ou "-" se ainda nao existir.
Saida (stdout): JSON com {"fila": [...], "total_pendente": N, "ja_feitos": N, "em_espera": N}.
--criar-corte N (so na 1a execucao, sem estado): poe no "_ignorar" todos os IDs fora das N ultimas linhas
da base e grava ./_estado.json. IDs em "_ignorar" nunca entram na fila.
Ordem: ID decrescente (prospect mais novo primeiro), depois o backlog.
Pulados por falta de fonte so voltam a fila depois de ESPERA_DIAS.
"""
import sys, re, json, argparse, datetime, pathlib, io

ESPERA_DIAS = 28

def linhas_xlsx(p):
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb["Inputs"] if "Inputs" in wb.sheetnames else wb.worksheets[0]
    cab = None
    for row in ws.iter_rows(values_only=True):
        vals = [str(v).strip() if v is not None else "" for v in row]
        if cab is None:
            if "ID" in vals and "Deal" in vals:
                cab = vals
            continue
        r = dict(zip(cab, vals))
        if re.fullmatch(r"\d+(\.0)?", r.get("ID", "")):
            r["ID"] = str(int(float(r["ID"])))
            yield r

def linhas_csv(p):
    import csv
    t = pathlib.Path(p).read_text(encoding="utf-8-sig")
    # delimitador: o que mais aparece na linha de cabecalho (virgula, ponto e virgula ou tab)
    linha_cab = next((l for l in t.splitlines() if "ID" in l and "Deal" in l), t.splitlines()[0])
    delim = max(",;\t", key=linha_cab.count)
    cab = None
    for cel in csv.reader(io.StringIO(t), delimiter=delim):
        cel = [c.strip() for c in cel]
        if cab is None:
            if "ID" in cel and "Deal" in cel:
                cab = cel
            continue
        r = dict(zip(cab, cel))
        if re.fullmatch(r"\d+(\.0)?", r.get("ID", "")):
            r["ID"] = str(int(float(r["ID"])))
            yield r

def linhas_txt(p):
    t = pathlib.Path(p).read_text(encoding="utf-8")
    if t.lstrip().startswith("["):
        t = json.loads(t)[0]["text"]
    cab = None
    for ln in t.split("\n"):
        cel = [c.strip() for c in ln.split("\t")]
        if cab is None:
            if "ID" in cel and "Deal" in cel:
                cab = cel
            continue
        if len(cel) > 3 and re.fullmatch(r"\d{2,7}", cel[cab.index("ID")] if cab.index("ID") < len(cel) else ""):
            yield dict(zip(cab, cel))

def main():
    sys.stdout.reconfigure(encoding="utf-8")  # Windows usa cp1252 por padrao e corrompe acentos
    ap = argparse.ArgumentParser()
    ap.add_argument("base"); ap.add_argument("pastas"); ap.add_argument("estado")
    ap.add_argument("--limite", type=int, default=25)
    ap.add_argument("--hoje", default=datetime.date.today().isoformat())
    ap.add_argument("--criar-corte", type=int)
    a = ap.parse_args()
    hoje = datetime.date.fromisoformat(a.hoje)

    if a.base.endswith((".xlsx", ".xlsm")):
        fonte = linhas_xlsx(a.base)
    elif a.base.endswith(".csv"):
        fonte = linhas_csv(a.base)
    else:
        fonte = linhas_txt(a.base)
    # so linhas de prospect de verdade (Status "1. Ativo", "2. Farming", "3. Inativo"...); descarta lixo de notas
    linhas = [r for r in fonte if re.match(r"\d\.\s", r.get("Status", ""))]
    base = {r["ID"]: r for r in linhas}
    feitos = {m.group(1) for n in json.loads(pathlib.Path(a.pastas).read_text(encoding="utf-8"))
              if (m := re.match(r"(\d+)\s*-", n))}
    estado = {} if a.estado == "-" else json.loads(pathlib.Path(a.estado).read_text(encoding="utf-8"))
    if a.criar_corte:
        manter = {r["ID"] for r in linhas[-a.criar_corte:]}
        estado["_ignorar"] = sorted((i for i in base if i not in manter), key=int)
        pathlib.Path("_estado.json").write_text(json.dumps(estado, ensure_ascii=False, indent=1), encoding="utf-8")
    ignorar = set(estado.get("_ignorar", []))

    fila, espera = [], 0
    # erros da revisao tem pasta, mas precisam rodar de novo
    feitos = feitos - {i for i, e in estado.items() if isinstance(e, dict) and e.get("status") == "erro"}
    for i in sorted((i for i in base if i not in feitos and i not in ignorar), key=int, reverse=True):
        e = estado.get(i)
        if e and e.get("status") == "pulado":
            if (hoje - datetime.date.fromisoformat(e["data"])).days < ESPERA_DIAS:
                espera += 1
                continue
        r = base[i]
        item = {k: r.get(k, "") for k in ("ID", "Deal", "Site", "Status", "Stage", "Macroproduto")}
        item["linha"] = {k: v for k, v in r.items() if k and v}  # linha inteira da base, para o subagente nao baixar a base de novo
        fila.append(item)
    print(json.dumps({"fila": fila[:a.limite], "total_pendente": len(fila),
                      "ja_feitos": len(feitos & set(base)), "em_espera": espera},
                     ensure_ascii=False, indent=1))

if __name__ == "__main__":
    main()
